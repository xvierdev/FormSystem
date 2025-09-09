# -*- coding: utf-8 -*-

# Imports existentes do seu projeto (todos mantidos)
import csv
import io
import pytz
import uuid
import logging
from collections import defaultdict
from datetime import datetime
from flask import Flask, render_template, request, redirect, url_for, session, flash, Response
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from sqlalchemy.orm import joinedload, selectinload
from openpyxl import Workbook
from openpyxl.styles import Font

# Novas importações, apenas para o QR Code
import socket
import base64
import qrcode

# --- Configurações Iniciais e Logger (sem alterações) ---
def setup_logger():
    for handler in logging.getLogger('summary').handlers[:]:
        logging.getLogger('summary').removeHandler(handler)
    logger = logging.getLogger('summary')
    logger.setLevel(logging.INFO)
    handler = logging.FileHandler('summary.log', encoding='utf-8')
    formatter = logging.Formatter(
        '%(asctime)s - %(message)s', datefmt='%d/%m/%Y %H:%M:%S')
    handler.setFormatter(formatter)
    logger.addHandler(handler)
    return logger

summary_logger = setup_logger()
app = Flask(__name__)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///questionario.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SECRET_KEY'] = 'uma_chave_secreta_muito_mais_forte_agora'
db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = 'login'
login_manager.login_message = "Você precisa estar logado para acessar esta página."
login_manager.login_message_category = "error"
SAO_PAULO_TZ = pytz.timezone('America/Sao_Paulo')

# --- Modelos de Dados (sem alterações) ---
class Admin(db.Model, UserMixin):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(256), nullable=False)
    def set_password(self, password): self.password_hash = generate_password_hash(password)
    def check_password(self, password): return check_password_hash(self.password_hash, password)

class Pergunta(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    texto = db.Column(db.String(500), nullable=False)
    alternativas = db.relationship('Alternativa', backref='pergunta', cascade="all, delete-orphan")
    resposta_correta = db.relationship('Alternativa', primaryjoin="and_(Pergunta.id==Alternativa.pergunta_id, Alternativa.correta==True)", uselist=False, viewonly=True, lazy='joined')

class Alternativa(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    texto = db.Column(db.String(200), nullable=False)
    correta = db.Column(db.Boolean, default=False, nullable=False)
    pergunta_id = db.Column(db.Integer, db.ForeignKey('pergunta.id'), nullable=False)

class Submissao(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    token = db.Column(db.String(36), unique=True, nullable=False, index=True)
    nome_aluno = db.Column(db.String(100), nullable=False, index=True)
    data_hora = db.Column(db.DateTime, default=datetime.utcnow)
    concluida = db.Column(db.Boolean, default=False, nullable=False, index=True)
    respostas = db.relationship('RespostaAluno', backref='submissao', cascade="all, delete-orphan")

class RespostaAluno(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    submissao_id = db.Column(db.Integer, db.ForeignKey('submissao.id'), nullable=False)
    pergunta_id = db.Column(db.Integer, db.ForeignKey('pergunta.id'), nullable=False)
    alternativa_escolhida_id = db.Column(db.Integer, db.ForeignKey('alternativa.id'), nullable=True)
    correta = db.Column(db.Boolean, nullable=False)
    pergunta = db.relationship('Pergunta', lazy='joined')
    alternativa_escolhida = db.relationship('Alternativa', lazy='joined')


# --- Funções de Apoio e Rotas ---

def get_local_ip():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(('10.255.255.255', 1))
        ip_address = s.getsockname()[0]
    except Exception:
        ip_address = '127.0.0.1'
    finally:
        s.close()
    return ip_address

@login_manager.user_loader
def load_user(user_id):
    return Admin.query.get(int(user_id))

def carregar_questionario_do_csv(file_stream):
    try:
        db.session.query(Alternativa).delete()
        db.session.query(Pergunta).delete()
        stream = io.TextIOWrapper(file_stream, encoding='utf-8')
        reader = csv.reader(stream, delimiter=';')
        perguntas_para_adicionar = []
        for row in reader:
            if not row or len(row) < 3: continue
            pergunta_texto = row[1]
            nova_pergunta = Pergunta(texto=pergunta_texto)
            alternativas = []
            for i in range(2, len(row)):
                alt_texto = row[i].strip()
                correta = '[correct]' in alt_texto
                if correta: alt_texto = alt_texto.replace('[correct]', '').strip()
                alternativas.append(Alternativa(texto=alt_texto, correta=correta))
            nova_pergunta.alternativas = alternativas
            perguntas_para_adicionar.append(nova_pergunta)
        if not perguntas_para_adicionar:
            raise ValueError("O arquivo CSV está vazio ou em formato inválido.")
        db.session.add_all(perguntas_para_adicionar)
        db.session.commit()
        num_perguntas = len(perguntas_para_adicionar)
        log_message = f"UPLOAD CSV SUCESSO - Admin: {current_user.username} - {num_perguntas} perguntas carregadas."
        summary_logger.info(log_message)
        return True, f"{num_perguntas} perguntas carregadas com sucesso!"
    except Exception as e:
        db.session.rollback()
        log_message = f"UPLOAD CSV FALHA - Admin: {current_user.username} - Erro: {e}"
        summary_logger.error(log_message)
        return False, f"Erro ao processar o arquivo: {e}"

@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated: return redirect(url_for('relatorio'))
    if request.method == 'POST':
        username, password = request.form.get('username'), request.form.get('password')
        user = Admin.query.filter_by(username=username).first()
        if user and user.check_password(password):
            login_user(user, remember=True)
            return redirect(url_for('relatorio'))
        else:
            flash('Usuário ou senha inválidos.', 'error')
    return render_template('login.html')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    flash('Você foi desconectado.', 'success')
    return redirect(url_for('login'))

@app.route('/upload', methods=['GET', 'POST'])
@login_required
def upload_csv():
    if request.method == 'POST':
        if 'csv_file' not in request.files:
            flash('Nenhum arquivo selecionado.', 'error')
            return redirect(request.url)
        file = request.files['csv_file']
        if file.filename == '':
            flash('Nenhum arquivo selecionado.', 'error')
            return redirect(request.url)
        if file and file.filename.endswith('.csv'):
            success, message = carregar_questionario_do_csv(file.stream)
            if success: flash(message, 'success')
            else: flash(message, 'error')
            return redirect(url_for('upload_csv'))
        else:
            flash('Por favor, envie um arquivo .csv válido.', 'error')
    return render_template('upload.html')

# =============================================================================
# ROTA MODIFICADA: /configuracoes (lógica do QR Code removida)
# =============================================================================
@app.route('/configuracoes', methods=['GET', 'POST'])
@login_required
def configuracoes():
    if request.method == 'POST':
        new_username = request.form.get('username').strip()
        new_password, confirm_password = request.form.get('password'), request.form.get('confirm_password')
        if not new_username:
            flash('O nome de usuário não pode estar em branco.', 'error')
            return redirect(url_for('configuracoes'))
        user = Admin.query.get(current_user.id)
        existing_user = Admin.query.filter(Admin.username == new_username, Admin.id != user.id).first()
        if existing_user:
            flash('Este nome de usuário já está em uso.', 'error')
            return redirect(url_for('configuracoes'))
        user.username = new_username
        if new_password:
            if new_password != confirm_password:
                flash('As senhas não coincidem.', 'error')
                return redirect(url_for('configuracoes'))
            user.set_password(new_password)
            flash('Usuário e senha atualizados com sucesso.', 'success')
        else:
            flash('Nome de usuário atualizado com sucesso.', 'success')
        db.session.commit()
        return redirect(url_for('configuracoes'))
    
    # A rota agora simplesmente renderiza o HTML sem a variável do QR code
    return render_template('configuracoes.html')

# =============================================================================
# ROTA MODIFICADA: / (index) (lógica do QR Code adicionada aqui)
# =============================================================================
@app.route('/')
def index():
    session.clear()

    # Gera o QR Code para a página inicial
    local_ip = get_local_ip()
    server_url = f"http://{local_ip}:5000"
    
    qr_img = qrcode.make(server_url)
    buffer = io.BytesIO()
    qr_img.save(buffer)
    buffer.seek(0)
    
    qr_code_image_b64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
    
    # Renderiza a página passando a imagem do QR Code
    return render_template('index.html', qr_code_image=qr_code_image_b64)

@app.route('/iniciar_questionario', methods=['POST'])
def iniciar_questionario():
    nome_aluno = request.form.get('nome_aluno', '').strip()
    if not nome_aluno:
        # Precisamos gerar o QR code aqui também para a página não quebrar em caso de erro
        local_ip = get_local_ip()
        server_url = f"http://{local_ip}:5000"
        qr_img = qrcode.make(server_url)
        buffer = io.BytesIO()
        qr_img.save(buffer)
        buffer.seek(0)
        qr_code_image_b64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
        return render_template('index.html', error="Por favor, digite seu nome.", qr_code_image=qr_code_image_b64)

    if Pergunta.query.count() == 0:
        return render_template('index.html', error="Nenhum questionário disponível no momento. Contate o administrador.")
    if Submissao.query.filter_by(nome_aluno=nome_aluno, concluida=True).first():
        return render_template('ja_submetido.html', nome_aluno=nome_aluno)
    token = str(uuid.uuid4())
    nova_submissao = Submissao(token=token, nome_aluno=nome_aluno)
    db.session.add(nova_submissao)
    db.session.commit()
    session['submission_token'] = token
    return redirect(url_for('questionario'))

@app.route('/questionario')
def questionario():
    token = session.get('submission_token')
    if not token: return redirect(url_for('index'))
    submissao = Submissao.query.filter_by(token=token, concluida=False).first()
    if not submissao: return redirect(url_for('index'))
    perguntas = Pergunta.query.options(joinedload(Pergunta.alternativas)).all()
    return render_template('questionario.html', perguntas=perguntas, nome_aluno=submissao.nome_aluno)

@app.route('/submeter_questionario', methods=['POST'])
def submeter_questionario():
    token = session.get('submission_token')
    if not token: return redirect(url_for('index'))
    submissao = Submissao.query.filter_by(token=token, concluida=False).first()
    if not submissao: return redirect(url_for('index'))
    perguntas_db = Pergunta.query.options(selectinload(Pergunta.alternativas)).all()
    alternativas_map = {alt.id: alt for p in perguntas_db for alt in p.alternativas}
    perguntas_map = {p.id: p for p in perguntas_db}
    total_acertos = 0
    respostas_para_salvar = []
    for pergunta_id, pergunta in perguntas_map.items():
        id_alternativa_str = request.form.get(f'pergunta_{pergunta_id}')
        alternativa_escolhida, correta = None, False
        if id_alternativa_str and id_alternativa_str.isdigit():
            id_alternativa = int(id_alternativa_str)
            if id_alternativa in alternativas_map:
                alternativa_escolhida = alternativas_map[id_alternativa]
                if alternativa_escolhida.correta:
                    correta = True
                    total_acertos += 1
        respostas_para_salvar.append(RespostaAluno(submissao_id=submissao.id, pergunta_id=pergunta.id, alternativa_escolhida_id=alternativa_escolhida.id if alternativa_escolhida else None, correta=correta))
    db.session.bulk_save_objects(respostas_para_salvar)
    submissao.concluida = True
    db.session.commit()
    total_perguntas = len(perguntas_map)
    log_message = f"SUBMISSÃO - Aluno: {submissao.nome_aluno} - Nota: {total_acertos}/{total_perguntas}"
    summary_logger.info(log_message)
    nome_aluno = submissao.nome_aluno
    session.clear()
    return render_template('agradecimento_com_nota.html', acertos=total_acertos, total=len(perguntas_map), nome_aluno=nome_aluno)

@app.route('/relatorio')
@login_required
def relatorio():
    submissoes = Submissao.query.filter_by(concluida=True).options(selectinload(Submissao.respostas).options(joinedload(RespostaAluno.pergunta).options(joinedload(Pergunta.resposta_correta)), joinedload(RespostaAluno.alternativa_escolhida))).order_by(Submissao.data_hora.desc()).all()
    total_perguntas_atual = Pergunta.query.count()
    return render_template('relatorio.html', submissoes=submissoes, total_perguntas_atual=total_perguntas_atual, SAO_PAULO_TZ=SAO_PAULO_TZ, utc=pytz.utc)

@app.route('/limpar_registros', methods=['POST'])
@login_required
def limpar_registros():
    try:
        num_respostas = db.session.query(RespostaAluno).delete()
        num_submissoes = db.session.query(Submissao).delete()
        db.session.commit()
        log_message = f"LIMPEZA DE DADOS - Admin: {current_user.username} - {num_submissoes} submissões e {num_respostas} respostas foram apagadas."
        summary_logger.info(log_message)
        flash(f'Todos os {num_submissoes} registros de alunos foram apagados com sucesso!', 'success')
    except Exception as e:
        db.session.rollback()
        summary_logger.error(f"LIMPEZA DE DADOS FALHA - Admin: {current_user.username} - Erro: {e}")
        flash(f'Ocorreu um erro ao limpar os registros: {e}', 'error')
    return redirect(url_for('relatorio'))

@app.route('/estatisticas')
@login_required
def estatisticas():
    perguntas = Pergunta.query.options(selectinload(Pergunta.alternativas)).all()
    respostas = RespostaAluno.query.all()
    charts_data = {}
    for p in perguntas:
        labels = [alt.texto for alt in p.alternativas]
        data = [0] * len(labels)
        charts_data[p.id] = {'texto': p.texto, 'labels': labels, 'data': data}
    for r in respostas:
        if r.pergunta_id in charts_data and r.alternativa_escolhida:
            try:
                idx = charts_data[r.pergunta_id]['labels'].index(r.alternativa_escolhida.texto)
                charts_data[r.pergunta_id]['data'][idx] += 1
            except ValueError: continue
    acertos_por_pergunta = defaultdict(int)
    erros_por_pergunta = defaultdict(int)
    total_respostas = len(respostas)
    for r in respostas:
        if r.correta: acertos_por_pergunta[r.pergunta_id] += 1
        else: erros_por_pergunta[r.pergunta_id] += 1
    mais_acertos = None
    if acertos_por_pergunta:
        id_mais_acertos = max(acertos_por_pergunta, key=acertos_por_pergunta.get)
        pergunta = db.session.get(Pergunta, id_mais_acertos)
        mais_acertos = {'texto': pergunta.texto, 'contagem': acertos_por_pergunta[id_mais_acertos]}
    mais_erros = None
    if erros_por_pergunta:
        id_mais_erros = max(erros_por_pergunta, key=erros_por_pergunta.get)
        pergunta = db.session.get(Pergunta, id_mais_erros)
        mais_erros = {'texto': pergunta.texto, 'contagem': erros_por_pergunta[id_mais_erros]}
    return render_template('estatisticas.html', charts=charts_data, mais_acertos=mais_acertos, mais_erros=mais_erros, total_respostas=total_respostas)

def get_todas_submissoes():
    return Submissao.query.filter_by(concluida=True).options(selectinload(Submissao.respostas).options(joinedload(RespostaAluno.pergunta).options(joinedload(Pergunta.resposta_correta)), joinedload(RespostaAluno.alternativa_escolhida))).order_by(Submissao.data_hora.desc()).all()

@app.route('/exportar/simplificado')
@login_required
def exportar_simplificado():
    submissoes = get_todas_submissoes()
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório Simplificado"
    headers = ["Data/Hora de Submissão", "Nome do Aluno", "Nota Final"]
    ws.append(headers)
    for cell in ws[1]: cell.font = Font(bold=True)
    for submissao in submissoes:
        data_hora = submissao.data_hora.replace(tzinfo=pytz.utc).astimezone(SAO_PAULO_TZ).strftime('%d/%m/%Y %H:%M:%S')
        acertos = sum(1 for r in submissao.respostas if r.correta)
        total = len(submissao.respostas)
        nota_final = f"{acertos}/{total}"
        ws.append([data_hora, submissao.nome_aluno, nota_final])
    mem_file = io.BytesIO()
    wb.save(mem_file)
    mem_file.seek(0)
    return Response(mem_file, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment;filename=relatorio_simplificado.xlsx'})

@app.route('/exportar/completo')
@login_required
def exportar_completo():
    submissoes = get_todas_submissoes()
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório Completo"
    headers = ["Nome do Aluno", "Pergunta", "Resposta do Aluno", "Resposta Correta", "Acertou?"]
    ws.append(headers)
    for cell in ws[1]: cell.font = Font(bold=True)
    for submissao in submissoes:
        for resposta in submissao.respostas:
            resposta_aluno = resposta.alternativa_escolhida.texto if resposta.alternativa_escolhida else "Não respondida"
            resposta_correta = resposta.pergunta.resposta_correta.texto if resposta.pergunta.resposta_correta else "N/A"
            acertou = "Sim" if resposta.correta else "Não"
            ws.append([submissao.nome_aluno, resposta.pergunta.texto, resposta_aluno, resposta_correta, acertou])
    mem_file = io.BytesIO()
    wb.save(mem_file)
    mem_file.seek(0)
    return Response(mem_file, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment;filename=relatorio_completo.xlsx'})

if __name__ == '__main__':
    with app.app_context():
        db.create_all()
        if not Admin.query.first():
            default_admin = Admin(username='admin')
            default_admin.set_password('admin')
            db.session.add(default_admin)
            db.session.commit()
            print("Usuário 'admin' com senha 'admin' criado.")
        if Pergunta.query.count() == 0:
            print("Nenhuma pergunta encontrada. Carregando 'mario.csv'...")
            try:
                with open('mario.csv', 'rb') as f:
                    carregar_questionario_do_csv(f)
            except FileNotFoundError:
                print("AVISO: 'mario.csv' não encontrado. O sistema iniciará sem perguntas.")
                print("Use a página de upload para carregar um questionário.")
    app.run(host='0.0.0.0', port=5000, debug=True)